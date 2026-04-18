"""Schedule 2 .xlsx generator — the Schedule of Exploratory Holes.

Produces a standalone workbook listing every exploratory hole in the
project: hole number, type, grid reference (placeholder until the
domain model carries coordinates), scheduled depth in metres, and
remarks (test selections and borehole instrumentation).

The Specification document's Chapter 2 embeds the same data as a Word
table; the row-building logic here and in the Specification generator
will be deduplicated into a shared helper once both feature branches
have merged.

Layout:
    Row 1: title  ("Schedule 2 — Exploratory Holes", bold, size 14)
    Row 3: bold column headers
    Row 4 onwards: one hole per row across columns A-E:
        A=Hole No., B=Type, C=Grid Reference, D=Scheduled Depth (m),
        E=Remarks

Scheduled depths are written as native numeric values (not
pre-formatted strings) so the workbook can be sorted or summed by Excel
directly. Depth cells may be empty when a hole type carries no
scheduled depth on the project (e.g. an unspecified trench depth).
"""

from collections.abc import Iterator
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from groundbill.models import Borehole, Project

_TITLE_ROW = 1
_HEADER_ROW = 3
_FIRST_DATA_ROW = 4

_HEADERS = ("Hole No.", "Type", "Grid Reference", "Scheduled Depth (m)", "Remarks")

_COLUMN_WIDTHS = {
    "A": 12.0,
    "B": 20.0,
    "C": 18.0,
    "D": 20.0,
    "E": 50.0,
}

_GRID_REFERENCE_TBC = "TBC"

_BOLD = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=14)


def generate_schedule_2(project: Project, output_path: Path | str) -> Path:
    """Write the Schedule 2 workbook for ``project``. Returns the absolute path."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule 2"
    _write_sheet(ws, project)

    path = Path(output_path)
    wb.save(path)
    return path.resolve()


def _write_sheet(ws: Worksheet, project: Project) -> None:
    for letter, width in _COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width

    title_cell = ws.cell(row=_TITLE_ROW, column=1, value="Schedule 2 \u2014 Exploratory Holes")
    title_cell.font = _TITLE_FONT

    for col_index, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=col_index, value=header)
        cell.font = _BOLD

    for row_offset, row in enumerate(_iter_rows(project)):
        r = _FIRST_DATA_ROW + row_offset
        for col_index, value in enumerate(row, start=1):
            ws.cell(row=r, column=col_index, value=value)


def _iter_rows(project: Project) -> Iterator[list]:
    """Yield one row of cell values per hole, in a stable type-by-type order.

    Each row is a 5-element list: [hole_number, hole_type, grid_reference,
    scheduled_depth, remarks]. Scheduled depth is a float (or None when
    the project has not yet assigned one to that hole); Excel will display
    an empty cell for None.
    """
    for bh in project.boreholes:
        yield _row(
            bh.hole_number,
            "Borehole",
            bh.total_schedule_depth_m,
            _borehole_remarks(bh),
        )
    for tp in project.trial_pits:
        yield _row(
            tp.trial_pit_number,
            "Trial Pit",
            tp.schedule_depth_m,
            _tests_remarks(tp.in_situ_tests),
        )
    for tr in project.trenches:
        yield _row(
            tr.trench_number,
            "Slit Trench",
            tr.overall_total_depth_m,  # may be None
            _tests_remarks(tr.in_situ_tests),
        )
    for ip in project.inspection_pits:
        yield _row(
            ip.inspection_pit_number,
            "Inspection Pit",
            ip.scheduled_depth_m,
            _tests_remarks(ip.in_situ_tests),
        )
    for ds in project.dynamic_samples:
        yield _row(
            ds.sample_number,
            "Dynamic Sample",
            ds.depth_m,
            _tests_remarks(ds.tests),
        )
    for sk in project.soakaways:
        yield _row(
            sk.soakaway_id,
            "Soakaway (BRE)",
            sk.schedule_depth_m,
            _tests_remarks(sk.in_situ_tests),
        )
    for dp in project.dynamic_probes:
        yield _row(dp.probe_number, "Dynamic Probe (DPH)", dp.depth_m, "")
    for cp in project.cpts:
        yield _row(
            cp.cpt_number,
            "CPT",
            cp.depth_m,
            "Piezocone" if cp.piezocone else "",
        )


def _row(hole_number: str, hole_type: str, scheduled_depth: float | None, remarks: str) -> list:
    return [hole_number, hole_type, _GRID_REFERENCE_TBC, scheduled_depth, remarks]


def _borehole_remarks(bh: Borehole) -> str:
    parts: list[str] = []
    if bh.tests:
        parts.append("Tests: " + ",".join(sorted(t.value for t in bh.tests)))
    if bh.piezometer_type.value != "none":
        parts.append(f"Instrumentation: {bh.piezometer_type.value}")
    return "; ".join(parts)


def _tests_remarks(tests: set) -> str:
    if not tests:
        return ""
    return "Tests: " + ",".join(sorted(t.value for t in tests))
