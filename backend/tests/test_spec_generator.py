"""Tests for the Specification .docx generator.

Byte-equality on a .docx is impossible (the zip embeds timestamps and
random IDs), so instead we re-open the generated file with python-docx
and assert on its content: cover-page text, chapter headings, table
structure, and route-not-supported errors.
"""

from pathlib import Path

import pytest
from docx import Document as load_docx

from groundbill.clauses import TBC_MARKER
from groundbill.generators import generate_spec
from groundbill.models import (
    AnticipatedGeology,
    ContractParties,
    ContractRoute,
    Project,
)
from tests.fixtures.basic_site import build_basic_site

# ---------------------------------------------------------------------------
# Cover page / file-level
# ---------------------------------------------------------------------------


def test_generated_file_exists_and_is_nonempty(tmp_path: Path) -> None:
    out = generate_spec(build_basic_site(), tmp_path / "spec.docx")
    assert out.exists()
    assert out.stat().st_size > 0


def test_cover_page_contains_project_name_and_address(tmp_path: Path) -> None:
    project = build_basic_site()
    out = generate_spec(project, tmp_path / "spec.docx")
    doc = load_docx(str(out))

    texts = [p.text for p in doc.paragraphs]
    assert any("Ground Investigation Specification" in t for t in texts), "cover title missing"
    assert project.name in texts, f"project name {project.name!r} not on cover page"
    assert project.site_address in texts, "site address not on cover page"


def test_cover_page_shows_contract_route_and_site_category(tmp_path: Path) -> None:
    project = build_basic_site()
    out = generate_spec(project, tmp_path / "spec.docx")
    doc = load_docx(str(out))
    texts = [p.text for p in doc.paragraphs]

    assert any(t.startswith("Contract route:") for t in texts)
    assert any(t.startswith("Site category:") for t in texts)
    assert any(t.startswith("Date:") for t in texts)


def test_preamble_clause_is_present(tmp_path: Path) -> None:
    project = build_basic_site()
    out = generate_spec(project, tmp_path / "spec.docx")
    doc = load_docx(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Preamble" in headings


# ---------------------------------------------------------------------------
# Chapter 1 — Introduction
# ---------------------------------------------------------------------------


def test_ch1_introduction_heading_and_subsections(tmp_path: Path) -> None:
    out = generate_spec(build_basic_site(), tmp_path / "spec.docx")
    doc = load_docx(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]

    assert "1 Introduction" in headings
    assert "1.1 Roles" in headings
    assert "1.2 Timescales" in headings
    assert "1.3 Anticipated Geology" in headings
    assert "1.3.1 Drift Geology" in headings
    assert "1.3.2 Solid Geology" in headings
    assert "1.3.3 Historical Ground Investigation Information" in headings
    assert "1.3.4 Mining Information" in headings


def test_ch1_roles_substitute_provided_parties(tmp_path: Path) -> None:
    project = build_basic_site().model_copy(
        update={
            "parties": ContractParties(
                employer_name="Test Employer Ltd.",
                psdp_organisation="Test Consulting",
                investigation_supervisor_organisation="Test IS Org",
            )
        }
    )
    out = generate_spec(project, tmp_path / "spec.docx")
    doc = load_docx(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)

    assert "Test Employer Ltd." in all_text
    assert "Test Consulting" in all_text
    assert "Test IS Org" in all_text


def test_ch1_missing_parties_render_as_tbc(tmp_path: Path) -> None:
    # basic_site fixture has no parties set — everything defaults
    out = generate_spec(build_basic_site(), tmp_path / "spec.docx")
    doc = load_docx(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)

    # The PSDP / Employer / IS lines all read "[TBC]"
    assert all_text.count(TBC_MARKER) >= 3


def test_ch1_anticipated_geology_uses_provided_text(tmp_path: Path) -> None:
    project = build_basic_site().model_copy(
        update={
            "anticipated_geology": AnticipatedGeology(
                drift_geology="Glacial till over fluvioglacial sands.",
                solid_geology="Carboniferous limestone.",
            )
        }
    )
    out = generate_spec(project, tmp_path / "spec.docx")
    doc = load_docx(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)

    assert "Glacial till over fluvioglacial sands." in all_text
    assert "Carboniferous limestone." in all_text
    # Historical and Mining were omitted — should render as TBC
    assert TBC_MARKER in all_text


# ---------------------------------------------------------------------------
# Chapter 2 — Ground Investigation Schedule
# ---------------------------------------------------------------------------


def test_ch2_heading_and_holes_table(tmp_path: Path) -> None:
    project = build_basic_site()
    out = generate_spec(project, tmp_path / "spec.docx")
    doc = load_docx(str(out))

    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "2 Ground Investigation Schedule" in headings

    assert len(doc.tables) == 1, "expected exactly one table — the holes table"
    table = doc.tables[0]

    expected_hole_rows = (
        len(project.boreholes)
        + len(project.trial_pits)
        + len(project.trenches)
        + len(project.inspection_pits)
        + len(project.dynamic_samples)
        + len(project.soakaways)
        + len(project.dynamic_probes)
        + len(project.cpts)
    )
    assert len(table.rows) == 1 + expected_hole_rows  # +1 for header

    header_cells = [cell.text for cell in table.rows[0].cells]
    assert header_cells == [
        "Hole No.",
        "Type",
        "Grid Reference",
        "Scheduled Depth (m)",
        "Remarks",
    ]

    data_column_0 = [row.cells[0].text for row in table.rows[1:]]
    assert "BH01" in data_column_0
    assert "TP01" in data_column_0
    assert "ST01" in data_column_0


def test_ch2_borehole_row_has_correct_type_and_depth(tmp_path: Path) -> None:
    out = generate_spec(build_basic_site(), tmp_path / "spec.docx")
    doc = load_docx(str(out))
    table = doc.tables[0]

    bh01_row = next(row for row in table.rows[1:] if row.cells[0].text == "BH01")
    assert bh01_row.cells[1].text == "Borehole"
    assert bh01_row.cells[3].text == "10.00"


def test_ch2_includes_drawing_reference_when_set(tmp_path: Path) -> None:
    project = build_basic_site().model_copy(
        update={"location_drawing_reference": "DWG-12345-SK001"}
    )
    out = generate_spec(project, tmp_path / "spec.docx")
    doc = load_docx(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)

    assert "DWG-12345-SK001" in all_text


def test_ch2_trench_commentary_appears_only_when_trenches_present(tmp_path: Path) -> None:
    with_trenches = build_basic_site()  # has 1 trench
    out1 = generate_spec(with_trenches, tmp_path / "with.docx")
    text_with = "\n".join(p.text for p in load_docx(str(out1)).paragraphs)

    without_trenches = with_trenches.model_copy(update={"trenches": []})
    out2 = generate_spec(without_trenches, tmp_path / "without.docx")
    text_without = "\n".join(p.text for p in load_docx(str(out2)).paragraphs)

    assert "slit trenches" in text_with.lower()
    assert "slit trenches" not in text_without.lower()


# ---------------------------------------------------------------------------
# Chapter 3 — Sampling Requirements
# ---------------------------------------------------------------------------


def test_ch3_heading_and_all_subsections_emitted(tmp_path: Path) -> None:
    """Sub-sections are always emitted so numbering is stable across projects."""
    out = generate_spec(build_basic_site(), tmp_path / "spec.docx")
    doc = load_docx(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]

    assert "3 Geotechnical Sampling and In-situ Testing Requirements" in headings
    assert "3.1 Hard Standing" in headings
    assert "3.2 Cable Percussion Boring" in headings
    assert "3.3 Rotary Drilling" in headings
    assert "3.4 Dynamic Probes" in headings
    assert "3.5 Trial Pits" in headings
    assert "3.6 Slit Trenching" in headings
    assert "3.7 In-situ Testing, Sampling and Monitoring" in headings
    assert "3.8 Contamination Avoidance and Aquifer Protection" in headings
    assert "3.9 Photography Requirements" in headings


def test_ch3_cable_percussion_renders_full_clause_when_used(tmp_path: Path) -> None:
    # basic_site has 3 CP boreholes
    out = generate_spec(build_basic_site(), tmp_path / "spec.docx")
    all_text = "\n".join(p.text for p in load_docx(str(out)).paragraphs)

    # Distinctive phrase from the CP boilerplate
    assert "Starting hole diameter shall be a minimum of 200 mm." in all_text


def test_ch3_cable_percussion_not_applicable_without_cp_method(tmp_path: Path) -> None:
    project = build_basic_site().model_copy(update={"boreholes": []})
    out = generate_spec(project, tmp_path / "spec.docx")
    doc = load_docx(str(out))

    # Walk paragraphs to find "3.2 Cable Percussion Boring" heading and check
    # the next non-empty paragraph is the Not-applicable marker.
    headings_and_bodies = [p for p in doc.paragraphs]
    idx = next(
        i for i, p in enumerate(headings_and_bodies) if p.text == "3.2 Cable Percussion Boring"
    )
    # First non-empty paragraph after the heading
    next_body = next(p.text for p in headings_and_bodies[idx + 1 :] if p.text.strip())
    assert next_body == "Not applicable to this contract."


def test_ch3_rotary_uses_live_rotary_borehole_count(tmp_path: Path) -> None:
    from groundbill.models import Borehole, DrillingMethod, DrillingPhase

    project = build_basic_site().model_copy(
        update={
            "boreholes": [
                Borehole(
                    hole_number="BH01",
                    phases=[
                        DrillingPhase(method=DrillingMethod.CABLE_PERCUSSION, depth_m=5.0),
                        DrillingPhase(method=DrillingMethod.ROTARY_CORE_HARD, depth_m=5.0),
                    ],
                    total_schedule_depth_m=10.0,
                ),
                Borehole(
                    hole_number="BH02",
                    phases=[
                        DrillingPhase(method=DrillingMethod.CABLE_PERCUSSION, depth_m=10.0),
                    ],
                    total_schedule_depth_m=10.0,
                ),
            ]
        }
    )
    out = generate_spec(project, tmp_path / "spec.docx")
    all_text = "\n".join(p.text for p in load_docx(str(out)).paragraphs)

    # Rotary section should be live and should mention "1 of the boreholes"
    assert "1 of the boreholes" in all_text


def test_ch3_dynamic_probes_uses_count_and_depth(tmp_path: Path) -> None:
    from groundbill.models import DynamicProbe

    project = build_basic_site().model_copy(
        update={
            "dynamic_probes": [
                DynamicProbe(probe_number=f"DPH0{i}", depth_m=6.0) for i in range(1, 4)
            ]
        }
    )
    out = generate_spec(project, tmp_path / "spec.docx")
    all_text = "\n".join(p.text for p in load_docx(str(out)).paragraphs)

    assert "3 no. Dynamic Probe Heavy" in all_text
    assert "6.0 m" in all_text


def test_ch3_trial_pits_uses_max_scheduled_depth(tmp_path: Path) -> None:
    from groundbill.models import TrialPit

    project = build_basic_site().model_copy(
        update={
            "trial_pits": [
                TrialPit(trial_pit_number="TP01", schedule_depth_m=3.0),
                TrialPit(trial_pit_number="TP02", schedule_depth_m=4.5),
            ]
        }
    )
    out = generate_spec(project, tmp_path / "spec.docx")
    all_text = "\n".join(p.text for p in load_docx(str(out)).paragraphs)

    assert "a scheduled depth of 4.5 m" in all_text


def test_ch3_hard_standing_conditional_on_paved_work(tmp_path: Path) -> None:
    from groundbill.models import TrialPit

    # basic_site fixture has only non-paved trial pits, so hard standing is N/A.
    out_no = generate_spec(build_basic_site(), tmp_path / "no.docx")
    text_no = "\n".join(p.text for p in load_docx(str(out_no)).paragraphs)

    # Find the 3.1 body
    doc_no = load_docx(str(out_no))
    paras = list(doc_no.paragraphs)
    idx = next(i for i, p in enumerate(paras) if p.text == "3.1 Hard Standing")
    next_body = next(p.text for p in paras[idx + 1 :] if p.text.strip())
    assert next_body == "Not applicable to this contract."

    # Now add a paved trial pit
    with_paved = build_basic_site().model_copy(
        update={
            "trial_pits": [
                TrialPit(trial_pit_number="TP01", schedule_depth_m=3.0, paved=True),
            ]
        }
    )
    out_yes = generate_spec(with_paved, tmp_path / "yes.docx")
    text_yes = "\n".join(p.text for p in load_docx(str(out_yes)).paragraphs)
    assert "concrete pavement of unknown thickness" in text_yes
    # Sanity: the boilerplate is not present in the N/A version
    assert "concrete pavement of unknown thickness" not in text_no


def test_ch3_insitu_testing_and_contamination_always_emitted(tmp_path: Path) -> None:
    """3.7 and 3.8 are always rendered with full boilerplate, even for a
    minimal project — every GI uses sampling and every site needs
    contamination-avoidance procedures."""
    project = build_basic_site().model_copy(
        update={
            "boreholes": [],
            "trial_pits": [],
            "trenches": [],
            "inspection_pits": [],
            "dynamic_samples": [],
            "soakaways": [],
            "dynamic_probes": [],
            "cpts": [],
        }
    )
    out = generate_spec(project, tmp_path / "spec.docx")
    all_text = "\n".join(p.text for p in load_docx(str(out)).paragraphs)

    assert "Standard Penetration Testing in accordance with IS EN ISO 22476-3" in all_text
    assert "Asbestos awareness procedures" in all_text


# ---------------------------------------------------------------------------
# Contract route gating
# ---------------------------------------------------------------------------


def test_pw_cf_route_is_not_supported_yet(tmp_path: Path) -> None:
    project = build_basic_site().model_copy(update={"contract_route": ContractRoute.PW_CF})
    with pytest.raises(NotImplementedError, match="PRIVATE"):
        generate_spec(project, tmp_path / "spec.docx")


# ---------------------------------------------------------------------------
# Regression: basic_site fixture still constructs cleanly after model changes
# ---------------------------------------------------------------------------


def test_basic_site_fixture_still_valid() -> None:
    project = build_basic_site()
    assert isinstance(project, Project)
    # New optional sub-models have their empty defaults
    assert project.parties.employer_name is None
    assert project.anticipated_geology.drift_geology is None
