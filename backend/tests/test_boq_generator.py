"""Tests for the BOQ .xlsx generator."""

from pathlib import Path

from openpyxl import load_workbook

from groundbill.generators import generate_boq
from groundbill.models import ContractRoute, Project, SiteCategory
from tests.fixtures.basic_site import build_basic_site


def _load(path: Path):
    wb = load_workbook(path, data_only=False)
    return wb, wb["Section A"]


def test_generated_file_has_expected_layout(tmp_path: Path):
    out = generate_boq(build_basic_site(), tmp_path / "boq.xlsx")
    assert out.exists()

    _, ws = _load(out)
    assert ws["A7"].value == "Number"
    assert ws["B7"].value == "Item description"
    assert ws["C7"].value == "Unit"
    assert ws["D7"].value == "Quantity"
    assert ws["E7"].value == "Rate"
    assert ws["F7"].value == "Amount"
    assert ws["A7"].font.bold is True

    assert ws["E8"].value == "€"
    assert ws["F8"].value == "€"

    assert ws["A9"].value == "A"
    assert ws["B9"].value == "General items, provisional services and additional items"
    assert ws["A9"].font.bold is True


def test_generated_items_start_at_row_10_with_correct_codes(tmp_path: Path):
    out = generate_boq(build_basic_site(), tmp_path / "boq.xlsx")
    _, ws = _load(out)
    assert ws["A10"].value == "A1"
    assert ws["C10"].value == "sum"
    assert ws["D10"].value == "Not Required"


def test_a8_quantity_is_written_as_integer(tmp_path: Path):
    out = generate_boq(build_basic_site(), tmp_path / "boq.xlsx")
    _, ws = _load(out)

    a8_row = _find_row_by_code(ws, "A8")
    assert ws.cell(row=a8_row, column=4).value == 9
    assert ws.cell(row=a8_row, column=3).value == "Nr"


def test_amount_cells_contain_qty_times_rate_formula(tmp_path: Path):
    out = generate_boq(build_basic_site(), tmp_path / "boq.xlsx")
    _, ws = _load(out)
    # First item row is 10; its Amount cell (F10) must be the formula.
    assert ws["F10"].value == '=IFERROR(D10*E10,"")'


def test_empty_project_produces_zero_for_a8(tmp_path: Path):
    project = Project(
        name="Empty",
        site_address="Nowhere",
        contract_route=ContractRoute.PRIVATE,
    )
    out = generate_boq(project, tmp_path / "boq.xlsx")
    _, ws = _load(out)
    a8_row = _find_row_by_code(ws, "A8")
    assert ws.cell(row=a8_row, column=4).value == 0


def test_yellow_site_emits_yellow_extra_over_items(tmp_path: Path):
    project = build_basic_site().model_copy(update={"site_category": SiteCategory.YELLOW})
    out = generate_boq(project, tmp_path / "boq.xlsx")
    _, ws = _load(out)
    codes = _all_codes(ws)
    assert "A3.1" in codes
    assert "A3.2" not in codes


def test_column_widths_match_contractor_boq_reference(tmp_path: Path):
    out = generate_boq(build_basic_site(), tmp_path / "boq.xlsx")
    _, ws = _load(out)
    assert ws.column_dimensions["A"].width == 9.2
    assert ws.column_dimensions["B"].width == 35.8
    assert ws.column_dimensions["C"].width == 10.2


def test_subheading_is_written_on_its_own_row_above_first_item(tmp_path: Path):
    out = generate_boq(build_basic_site(), tmp_path / "boq.xlsx")
    ws = load_workbook(out)["Section G"]

    assert ws["A9"].value == "G"
    assert ws["B9"].value == "Geophysical testing"

    # Row 10 is the sub-heading (column B only); the first item follows on row 11.
    assert ws["A10"].value is None
    assert ws["B10"].value == "Land-based mapping techniques"
    assert ws["B10"].font.bold is True
    assert ws["B10"].font.underline == "single"
    assert ws["A11"].value == "G1"
    assert ws["C11"].value == "m²"
    assert ws["F11"].value == '=IFERROR(D11*E11,"")'

    # Same row positions as the reference workbook: G6 on row 17, G10 on row 22.
    assert ws["B16"].value == "Borehole geophysical surveying"
    assert ws["A17"].value == "G6"
    assert ws["A22"].value == "G10"


def _find_row_by_code(ws, code: str) -> int:
    for r in range(10, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == code:
            return r
    raise AssertionError(f"code {code!r} not found in column A")


def _all_codes(ws) -> list[str]:
    codes = []
    for r in range(10, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None:
            codes.append(v)
    return codes
