"""Tests for the standalone Schedule 2 .xlsx generator."""

from pathlib import Path

from openpyxl import load_workbook

from groundbill.generators import generate_schedule_2
from groundbill.models import (
    CPT,
    Borehole,
    ContractRoute,
    DrillingMethod,
    DrillingPhase,
    DynamicProbe,
    DynamicSample,
    InspectionPit,
    Project,
    Soakaway,
    Trench,
    TrialPit,
)
from tests.fixtures.basic_site import build_basic_site


def _load(path: Path):
    wb = load_workbook(path, data_only=False)
    return wb, wb["Schedule 2"]


def test_generated_file_exists_and_opens(tmp_path: Path) -> None:
    out = generate_schedule_2(build_basic_site(), tmp_path / "schedule_2.xlsx")
    assert out.exists()
    assert out.stat().st_size > 0
    wb, ws = _load(out)
    assert ws.title == "Schedule 2"


def test_title_and_header_row_layout(tmp_path: Path) -> None:
    out = generate_schedule_2(build_basic_site(), tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)

    # Title on row 1, column A
    assert "Schedule 2" in ws["A1"].value
    assert "Exploratory Holes" in ws["A1"].value
    assert ws["A1"].font.bold is True

    # Headers on row 3, columns A-E
    assert ws["A3"].value == "Hole No."
    assert ws["B3"].value == "Type"
    assert ws["C3"].value == "Grid Reference"
    assert ws["D3"].value == "Scheduled Depth (m)"
    assert ws["E3"].value == "Remarks"
    assert ws["A3"].font.bold is True


def test_basic_site_produces_row_per_hole_in_type_order(tmp_path: Path) -> None:
    project = build_basic_site()
    out = generate_schedule_2(project, tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)

    expected_count = (
        len(project.boreholes)
        + len(project.trial_pits)
        + len(project.trenches)
        + len(project.inspection_pits)
        + len(project.dynamic_samples)
        + len(project.soakaways)
        + len(project.dynamic_probes)
        + len(project.cpts)
    )
    # Row 4 is first data row; last data row index is 4 + count - 1
    # Find last non-empty hole number cell.
    actual_numbers = []
    for r in range(4, 4 + expected_count + 2):
        v = ws.cell(row=r, column=1).value
        if v is None:
            break
        actual_numbers.append(v)
    assert len(actual_numbers) == expected_count

    # basic_site ordering: 3 BHs, 2 TPs, 1 Trench, 1 CPT, 1 DS — but
    # in the generator's row order boreholes -> TP -> trench -> inspection
    # pits -> dynamic samples -> soakaways -> dynamic probes -> CPT. So:
    assert actual_numbers[0] == "BH01"
    assert actual_numbers[1] == "BH02"
    assert actual_numbers[2] == "BH03"
    assert actual_numbers[3] == "TP01"
    assert actual_numbers[5] == "ST01"
    # Dynamic samples come before CPT in the row order
    assert "CPT01" in actual_numbers
    assert "DS01" in actual_numbers


def test_borehole_row_has_numeric_depth_and_correct_type(tmp_path: Path) -> None:
    out = generate_schedule_2(build_basic_site(), tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)

    # BH01 is on row 4 (first borehole)
    assert ws["A4"].value == "BH01"
    assert ws["B4"].value == "Borehole"
    assert ws["C4"].value == "TBC"
    # Depth is stored as a number (not a string), so it's sortable in Excel.
    # openpyxl round-trips whole-number floats as int, so accept either.
    assert ws["D4"].value == 10
    assert isinstance(ws["D4"].value, int | float)
    assert not isinstance(ws["D4"].value, str)


def test_trial_pit_remarks_include_tests(tmp_path: Path) -> None:
    project = build_basic_site()
    out = generate_schedule_2(project, tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)

    tp_row = _find_row_by_number(ws, "TP01")
    assert ws.cell(row=tp_row, column=2).value == "Trial Pit"
    remarks = ws.cell(row=tp_row, column=5).value
    # basic_site sets {InSituTest.DCP} on the trial pits
    assert remarks is not None
    assert "DCP" in remarks


def test_cpt_piezocone_flag_shown_in_remarks(tmp_path: Path) -> None:
    project = Project(
        name="CPT only",
        site_address="Somewhere",
        contract_route=ContractRoute.PRIVATE,
        cpts=[
            CPT(cpt_number="CPT01", depth_m=15.0, piezocone=True),
            CPT(cpt_number="CPT02", depth_m=15.0, piezocone=False),
        ],
    )
    out = generate_schedule_2(project, tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)

    r1 = _find_row_by_number(ws, "CPT01")
    r2 = _find_row_by_number(ws, "CPT02")
    assert ws.cell(row=r1, column=5).value == "Piezocone"
    assert ws.cell(row=r2, column=5).value in ("", None)


def test_trench_with_no_overall_depth_writes_empty_cell(tmp_path: Path) -> None:
    project = Project(
        name="Trench only",
        site_address="Somewhere",
        contract_route=ContractRoute.PRIVATE,
        trenches=[Trench(trench_number="ST01")],
    )
    out = generate_schedule_2(project, tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)

    r = _find_row_by_number(ws, "ST01")
    assert ws.cell(row=r, column=2).value == "Slit Trench"
    assert ws.cell(row=r, column=4).value is None


def test_borehole_with_piezometer_shows_instrumentation_in_remarks(tmp_path: Path) -> None:
    project = Project(
        name="BH with PIE",
        site_address="Somewhere",
        contract_route=ContractRoute.PRIVATE,
        boreholes=[
            Borehole(
                hole_number="BH10",
                phases=[DrillingPhase(method=DrillingMethod.CABLE_PERCUSSION, depth_m=10.0)],
                total_schedule_depth_m=10.0,
                piezometer=True,
            )
        ],
    )
    out = generate_schedule_2(project, tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)
    r = _find_row_by_number(ws, "BH10")
    assert ws.cell(row=r, column=5).value == "Instrumentation: piezometer"


def test_borehole_with_piezometer_and_standpipe_lists_both(tmp_path: Path) -> None:
    project = Project(
        name="BH with PIE and SP",
        site_address="Somewhere",
        contract_route=ContractRoute.PRIVATE,
        boreholes=[
            Borehole(
                hole_number="BH11",
                phases=[DrillingPhase(method=DrillingMethod.CABLE_PERCUSSION, depth_m=10.0)],
                total_schedule_depth_m=10.0,
                piezometer=True,
                standpipe=True,
            )
        ],
    )
    out = generate_schedule_2(project, tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)
    r = _find_row_by_number(ws, "BH11")
    assert ws.cell(row=r, column=5).value == "Instrumentation: piezometer, standpipe"


def test_empty_project_produces_headers_only(tmp_path: Path) -> None:
    project = Project(
        name="Empty",
        site_address="Nowhere",
        contract_route=ContractRoute.PRIVATE,
    )
    out = generate_schedule_2(project, tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)

    # Header still written
    assert ws["A3"].value == "Hole No."
    # Row 4 (first data row) is empty
    assert ws["A4"].value is None


def test_all_hole_types_render_with_depth(tmp_path: Path) -> None:
    project = Project(
        name="Kitchen sink",
        site_address="Somewhere",
        contract_route=ContractRoute.PRIVATE,
        boreholes=[
            Borehole(
                hole_number="BH01",
                phases=[DrillingPhase(method=DrillingMethod.CABLE_PERCUSSION, depth_m=12.0)],
                total_schedule_depth_m=12.0,
            )
        ],
        trial_pits=[TrialPit(trial_pit_number="TP01", schedule_depth_m=3.5)],
        trenches=[Trench(trench_number="ST01", overall_total_depth_m=2.0)],
        inspection_pits=[InspectionPit(inspection_pit_number="IP01", scheduled_depth_m=1.2)],
        dynamic_samples=[DynamicSample(sample_number="DS01", depth_m=6.0)],
        soakaways=[Soakaway(soakaway_id="SK01", schedule_depth_m=2.5)],
        dynamic_probes=[DynamicProbe(probe_number="DPH01", depth_m=8.0)],
        cpts=[CPT(cpt_number="CPT01", depth_m=20.0)],
    )
    out = generate_schedule_2(project, tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)

    # Collect {hole_number: (type, depth)}
    rows: dict[str, tuple[str, float]] = {}
    for r in range(4, 20):
        n = ws.cell(row=r, column=1).value
        if n is None:
            break
        rows[n] = (ws.cell(row=r, column=2).value, ws.cell(row=r, column=4).value)

    assert rows["BH01"] == ("Borehole", 12.0)
    assert rows["TP01"] == ("Trial Pit", 3.5)
    assert rows["ST01"] == ("Slit Trench", 2.0)
    assert rows["IP01"] == ("Inspection Pit", 1.2)
    assert rows["DS01"] == ("Dynamic Sample", 6.0)
    assert rows["SK01"] == ("Soakaway (BRE)", 2.5)
    assert rows["DPH01"] == ("Dynamic Probe (DPH)", 8.0)
    assert rows["CPT01"] == ("CPT", 20.0)


def test_column_widths_applied(tmp_path: Path) -> None:
    out = generate_schedule_2(build_basic_site(), tmp_path / "schedule_2.xlsx")
    _, ws = _load(out)
    assert ws.column_dimensions["A"].width == 12.0
    assert ws.column_dimensions["B"].width == 20.0
    assert ws.column_dimensions["E"].width == 50.0


def _find_row_by_number(ws, hole_number: str) -> int:
    for r in range(4, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == hole_number:
            return r
    raise AssertionError(f"hole number {hole_number!r} not found in column A")
