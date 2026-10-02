"""Regression guard: engine output must match the reference workbook row for row.

For each re-translated section this reads the corresponding sheet of
`reference/excel/4_BOQ_Contractor_Rev_A.xlsx` (the issued Contractor BOQ) and
asserts that the engine emits the same item codes, descriptions, units,
sub-headings and footnotes, in the same order.

All twelve sections, A to L, are covered.
"""

import re
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import load_workbook

from groundbill.engine import (
    BoqItem,
    compute_section_a,
    compute_section_b,
    compute_section_c,
    compute_section_d,
    compute_section_e,
    compute_section_f,
    compute_section_g,
    compute_section_h,
    compute_section_i,
    compute_section_j,
    compute_section_k,
    compute_section_l,
)
from groundbill.models import ContractRoute, Project

_CONTRACTOR_BOQ = (
    Path(__file__).resolve().parents[2] / "reference" / "excel" / "4_BOQ_Contractor_Rev_A.xlsx"
)

# Row 9 of every sheet is the section title; item and sub-heading rows follow.
_FIRST_BODY_ROW = 10

_VERIFIED_SECTIONS: dict[str, Callable[[Project], list[BoqItem]]] = {
    "A": compute_section_a,
    "B": compute_section_b,
    "C": compute_section_c,
    "D": compute_section_d,
    "E": compute_section_e,
    "F": compute_section_f,
    "G": compute_section_g,
    "H": compute_section_h,
    "I": compute_section_i,
    "J": compute_section_j,
    "K": compute_section_k,
    "L": compute_section_l,
}


def _norm(value: object) -> str:
    """Collapse runs of whitespace and trim, so stray trailing spaces in cells don't matter."""
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


@pytest.fixture(scope="module")  # module scope: open the workbook once for all tests in this file
def contractor_boq():
    return load_workbook(_CONTRACTOR_BOQ)


_Row = tuple[str, str, str, str | None, str | None, str | None]


def _workbook_rows(ws) -> list[_Row]:
    """Return (code, description, unit, sub-heading, its code, note) for each item row.

    An item row has a code in column A and a unit in column C. A sub-heading
    row has text in column B and no unit; it is attached to the next item row.
    In Section K the sub-heading rows also carry a code in column A. A footnote
    row has text beginning "Note" or "(Note" in column B; it is attached to the
    item row above it.
    """
    rows: list[_Row] = []
    pending_subheading: str | None = None
    pending_code: str | None = None
    for r in range(_FIRST_BODY_ROW, ws.max_row + 1):
        code, desc, unit = (_norm(ws.cell(r, c).value) for c in (1, 2, 3))
        if code and unit:
            rows.append((code, desc, unit, pending_subheading, pending_code, None))
            pending_subheading = pending_code = None
        elif desc:
            if desc.lower().startswith("total "):
                break
            if re.match(r"\(?note\b", desc, re.IGNORECASE) and not code:
                rows[-1] = (*rows[-1][:5], desc)
                continue
            pending_subheading = desc
            pending_code = code or None
    return rows


@pytest.mark.parametrize("letter", sorted(_VERIFIED_SECTIONS))
def test_engine_matches_contractor_workbook(contractor_boq, letter: str):
    expected = _workbook_rows(contractor_boq[f"Section {letter}"])
    project = Project(name="Empty", site_address="Nowhere", contract_route=ContractRoute.PRIVATE)
    actual = [
        (
            i.code,
            _norm(i.description),
            _norm(i.unit),
            i.subheading,
            i.subheading_code,
            _norm(i.note) or None,
        )
        for i in _VERIFIED_SECTIONS[letter](project)
    ]
    assert actual == expected
